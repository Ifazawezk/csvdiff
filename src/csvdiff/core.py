import csv
import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

Row = dict[str, str]
Key = tuple[str, ...]


class DuplicateKeyError(ValueError):
    pass


@dataclass
class Change:
    key: Key
    # column -> (old value, new value)
    cells: dict[str, tuple[str, str]]


@dataclass
class CompareOptions:
    ignore_columns: tuple[str, ...] = ()
    # Numbers within this absolute difference count as equal (e.g. 0.01 for rounding).
    tolerance: float = 0.0
    ignore_case: bool = False
    # Ignore leading/trailing spaces and repeated inner spaces.
    ignore_whitespace: bool = False

    def same(self, a: str, b: str) -> bool:
        if a == b:
            return True
        if self.ignore_whitespace:
            a, b = " ".join(a.split()), " ".join(b.split())
        if self.ignore_case:
            a, b = a.casefold(), b.casefold()
        if a == b:
            return True
        if self.tolerance:
            x, y = _to_number(a), _to_number(b)
            if x is not None and y is not None:
                return abs(x - y) <= self.tolerance + 1e-12
        return False


def _to_number(s: str) -> float | None:
    try:
        return float(s.strip().replace(",", ""))
    except ValueError:
        return None


@dataclass
class Diff:
    key_columns: list[str]
    added: list[Row] = field(default_factory=list)
    removed: list[Row] = field(default_factory=list)
    changed: list[Change] = field(default_factory=list)
    columns_added: list[str] = field(default_factory=list)
    columns_removed: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not (self.added or self.removed or self.changed or self.columns_added or self.columns_removed)


def load_xlsx(path: str | Path, sheet: str | None = None) -> tuple[list[str], list[Row]]:
    """Read one sheet of an Excel file into (header, rows), with every value as text."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("reading .xlsx files needs openpyxl: pip install openpyxl") from None
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    first = next(rows_iter, None) or ()
    header = [str(h).strip() if h is not None else "" for h in first]
    rows = []
    for values in rows_iter:
        if values is None or all(v is None for v in values):
            continue
        rows.append({h: _cell_text(v) for h, v in zip(header, values)})
    wb.close()
    return header, rows


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 42 as 42.0
    if isinstance(value, dt.datetime):
        # Excel has no separate date type, so midnight usually means "just a date".
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def load_table(path: str | Path, encoding: str = "utf-8-sig", delimiter: str | None = None,
               sheet: str | None = None) -> tuple[list[str], list[Row]]:
    """Read a .csv/.tsv/.txt or .xlsx file."""
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        return load_xlsx(path, sheet)
    return load_csv(path, encoding, delimiter)


def load_csv(path: str | Path, encoding: str = "utf-8-sig", delimiter: str | None = None) -> tuple[list[str], list[Row]]:
    """Read a CSV file into (header, rows). The delimiter is sniffed if not given."""
    with open(path, newline="", encoding=encoding) as f:
        sample = f.read(64 * 1024)
        f.seek(0)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.DictReader(f, delimiter=delimiter)
        # Exports from Excel and some databases often pad header names with spaces.
        header = [h.strip() for h in reader.fieldnames or []]
        reader.fieldnames = header
        rows = [dict(r) for r in reader]
        return header, rows


def _index(rows: list[Row], key_columns: list[str], label: str) -> dict[Key, Row]:
    index: dict[Key, Row] = {}
    for n, row in enumerate(rows, start=2):  # line 1 is the header
        missing = [k for k in key_columns if k not in row]
        if missing:
            raise KeyError(f"key column(s) {missing} not found in {label}")
        key = tuple(row[k] for k in key_columns)
        if key in index:
            raise DuplicateKeyError(f"duplicate key {key} in {label} (line {n})")
        index[key] = row
    return index


def _columns(rows: list[Row], header: list[str] | None, key_columns: list[str], label: str) -> list[str]:
    """A file's columns: its header if we have one, otherwise the keys of its first row."""
    if header is None:
        return list(rows[0]) if rows else []
    missing = [k for k in key_columns if k not in header]
    if missing:
        raise KeyError(f"key column(s) {missing} not found in {label}")
    return list(header)


def diff_rows(old: list[Row], new: list[Row], key_columns: list[str],
              options: CompareOptions | None = None,
              old_columns: list[str] | None = None, new_columns: list[str] | None = None) -> Diff:
    """Match rows by key_columns and compare the columns both files share.

    Pass the headers from load_table as old_columns and new_columns. Without them the
    columns are read from the first row, so a table with no rows appears to have none.
    """
    opts = options or CompareOptions()
    old_cols = _columns(old, old_columns, key_columns, "old file")
    new_cols = _columns(new, new_columns, key_columns, "new file")
    old_idx = _index(old, key_columns, "old file")
    new_idx = _index(new, key_columns, "new file")

    shared = [c for c in old_cols if c in new_cols and c not in key_columns and c not in opts.ignore_columns]

    d = Diff(
        key_columns=key_columns,
        columns_added=[c for c in new_cols if c not in old_cols],
        columns_removed=[c for c in old_cols if c not in new_cols],
    )
    for key, row in new_idx.items():
        if key not in old_idx:
            d.added.append(row)
    for key, row in old_idx.items():
        if key not in new_idx:
            d.removed.append(row)
            continue
        other = new_idx[key]
        cells = {c: (row[c], other[c]) for c in shared if not opts.same(row[c], other[c])}
        if cells:
            d.changed.append(Change(key, cells))
    return d
